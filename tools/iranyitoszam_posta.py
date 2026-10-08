#!/usr/bin/env python3
"""
A Magyar Posta irányítószám-táblázatából (Iranyitoszam-Internet_uj.xlsx)
elkészíti a plugin data/iranyitoszam.csv.gz fájlját.

Használat:  python3 -I tools/iranyitoszam_posta.py <posta.xlsx> data/iranyitoszam.csv.gz
Kell hozzá: openpyxl (pip install openpyxl).

Mit csinál:
- „Települések” lap: IRSZ, Település, Településrész – ahogy van (a szóközöket levágja).
- A nagyvárosok (Budapest, Miskolc, Debrecen, Szeged, Pécs, Győr) kódjai a
  Posta fájlban külön utcajegyzék-lapon vannak; innen vesszük fel azokat a
  kódokat, amelyek a „Települések” lapon nincsenek. Budapestnél a kerület a
  kód 2–3. számjegyéből jön (1051 → V. kerület).
- Kimenet: „irsz;telepules;telepulesresz” fejléc, irányítószám, majd név
  szerint rendezve, gzip (mtime 0, így azonos bemenetből bájtra azonos fájl).
"""
import gzip
import io
import sys

import openpyxl

VAROSI_LAPOK = {
    'Bp.u.': 'Budapest',
    'Miskolc u.': 'Miskolc',
    'Debrecen u.': 'Debrecen',
    'Szeged u.': 'Szeged',
    'Pécs u.': 'Pécs',
    'Győr u.': 'Győr',
}

ROMAI = ['', 'I.', 'II.', 'III.', 'IV.', 'V.', 'VI.', 'VII.', 'VIII.', 'IX.', 'X.', 'XI.', 'XII.',
         'XIII.', 'XIV.', 'XV.', 'XVI.', 'XVII.', 'XVIII.', 'XIX.', 'XX.', 'XXI.', 'XXII.', 'XXIII.']


def tiszta(ertek):
    return ' '.join(str(ertek).split()) if ertek is not None else ''


def main(be, ki):
    wb = openpyxl.load_workbook(be, read_only=True, data_only=True)
    sorok = set()

    for sor in wb['Települések'].iter_rows(values_only=True):
        if not sor or not isinstance(sor[0], int):
            continue
        irsz, nev, resz = sor[0], tiszta(sor[1]), tiszta(sor[2] if len(sor) > 2 else '')
        if 1000 <= irsz <= 9999 and nev:
            sorok.add((irsz, nev, resz))

    megvan = {(i, n) for i, n, _ in sorok}

    for lap, varos in VAROSI_LAPOK.items():
        for sor in wb[lap].iter_rows(values_only=True):
            if not sor or not isinstance(sor[0], int) or not 1000 <= sor[0] <= 9999:
                continue
            irsz = sor[0]
            if (irsz, varos) in megvan:
                continue
            resz = ''
            if varos == 'Budapest':
                ker = (irsz // 10) % 100
                resz = ROMAI[ker] + ' kerület' if 1 <= ker <= 23 else ''
            sorok.add((irsz, varos, resz))
            megvan.add((irsz, varos))

    kimenet = io.StringIO()
    kimenet.write('irsz;telepules;telepulesresz\n')
    for irsz, nev, resz in sorted(sorok, key=lambda s: (s[0], s[1], s[2])):
        kimenet.write('%04d;%s;%s\n' % (irsz, nev.replace(';', ','), resz.replace(';', ',')))

    with open(ki, 'wb') as f:
        with gzip.GzipFile(fileobj=f, mode='wb', mtime=0, filename='') as gz:
            gz.write(kimenet.getvalue().encode('utf-8'))

    print('%d sor, %d település, %d irányítószám, %d településrészes sor' % (
        len(sorok), len({s[1] for s in sorok}), len({s[0] for s in sorok}), sum(1 for s in sorok if s[2])))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
